"""Genera Modelo_demanda_oferta_IoT_LosRios.xlsx con fórmulas vivas.

Las cifras observadas (SRI, INEC, SENESCYT vía MINTEL, UTEQ) van en hojas de datos.
Los supuestos (celdas amarillas) son de trabajo y no tienen fuente oficial: deben
sustituirse con la encuesta a empleadores. Todo lo demás se calcula con fórmulas.
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
DATOS = os.path.join(AQUI, '..', 'datos')
SALIDA = os.path.join(AQUI, '..', 'Modelo_demanda_oferta_IoT_LosRios.xlsx')

AZUL = '1F4E79'
AMARILLO = PatternFill('solid', fgColor='FFF2CC')
GRIS = PatternFill('solid', fgColor='F2F2F2')
CAB = PatternFill('solid', fgColor=AZUL)
FUENTE_CAB = Font(bold=True, color='FFFFFF', name='Calibri')
NEGRITA = Font(bold=True, name='Calibri')
fino = Side(style='thin', color='BFBFBF')
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)


def cab(ws, fila, textos, col=1):
    for i, t in enumerate(textos):
        c = ws.cell(row=fila, column=col + i, value=t)
        c.fill, c.font, c.border = CAB, FUENTE_CAB, BORDE
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')


def celda(ws, fila, col, valor, fmt=None, entrada=False, negrita=False):
    c = ws.cell(row=fila, column=col, value=valor)
    c.border = BORDE
    if fmt:
        c.number_format = fmt
    if entrada:
        c.fill = AMARILLO
    if negrita:
        c.font = NEGRITA
    return c


def titulo(ws, texto, sub=None):
    ws['A1'] = texto
    ws['A1'].font = Font(bold=True, size=14, color=AZUL, name='Calibri')
    if sub:
        ws['A2'] = sub
        ws['A2'].font = Font(italic=True, size=10, name='Calibri')


def anchos(ws, lista):
    for i, w in enumerate(lista, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def leer_csv(nombre):
    with open(os.path.join(DATOS, nombre), encoding='utf-8', newline='') as f:
        return list(csv.reader(f))


wb = Workbook()

# ------------------------------------------------------------------ Léame
ws = wb.active
ws.title = 'Léame'
titulo(ws, 'Modelo de demanda y oferta de profesionales IoT en Los Ríos',
       'Borrador de trabajo para el proyecto de carrera de Ingeniería en Internet de las Cosas (UTEQ). Consulta de datos: 26-IX-2026.')
lineas = [
    ('Qué es', 'Un ejercicio de dimensionamiento con tres escenarios. Cruza empresas formales de Los Ríos (SRI e INEC) con supuestos de adopción y contratación de perfiles IoT, y lo compara con la oferta de graduados TIC.'),
    ('Qué NO es', 'No es una proyección estadística. Ninguna fuente oficial publica la adopción de IoT por sector en Ecuador; esos porcentajes son supuestos de trabajo. Deben sustituirse con la encuesta a empleadores antes de presentar el proyecto al CES.'),
    ('Código de color', 'Celdas amarillas = supuestos que usted puede cambiar. Celdas blancas = datos observados o fórmulas. Todo el modelo se recalcula al editar una celda amarilla.'),
    ('Hojas', 'SRI_LosRios y INEC_LosRios: datos observados. Supuestos: entradas. Demanda, Oferta, Brecha y Sensibilidad: cálculos.'),
    ('Horizonte', 'Año base 2026. "4 años" = 2030 y "5 años" = 2031. Con ingreso en 2027 y 8 períodos, la primera promoción de la carrera se gradúa en 2031.'),
    ('Fuentes', 'SRI: catastro RUC de Los Ríos (datos abiertos, actualizado el 1-IX-2026). INEC: Registro Estadístico de Empresas 2025, ENEMDU anual 2025, Censo 2022. SENESCYT: graduados TIC 2020 y 2021 tomados de la Política de Transformación Digital 2025-2030 (MINTEL). UTEQ: informe de rendición de cuentas 2024.'),
    ('Límites de los datos', 'El catastro RUC no trae tamaño de empresa; se usa "sociedad o persona natural obligada a llevar contabilidad" como aproximación de empresa formal. No se obtuvo de la SENESCYT la serie de graduados TIC por provincia (el servidor de sus datos abiertos no respondió); la oferta se estima con proxies que se detallan en la hoja Oferta.'),
    ('Privacidad', 'El catastro RUC incluye nombres de personas naturales. Este libro solo contiene conteos agregados.'),
]
for i, (k, v) in enumerate(lineas, start=4):
    ws.cell(row=i, column=1, value=k).font = NEGRITA
    c = ws.cell(row=i, column=2, value=v)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[i].height = 62
anchos(ws, [22, 120])

# ------------------------------------------------------------------ SRI
ws = wb.create_sheet('SRI_LosRios')
titulo(ws, 'Contribuyentes activos de Los Ríos (SRI, catastro RUC)',
       'Fuente: SRI, datos abiertos, "Registro Único de Contribuyentes / Los Ríos", actualizado el 1-IX-2026. Un RUC se cuenta una vez, con la CIIU de su establecimiento matriz.')
cab(ws, 4, ['Grupo de actividad (CIIU Rev. 4.1)', 'Sociedades activas', 'Personas naturales obligadas a contabilidad', 'Empresas formales', 'Total RUC activos (incluye no obligados)'])
grupos = leer_csv('sri_grupos.csv')[1:]
fila = 5
FILAS_SRI = {}
for g in grupos:
    nombre, soc, pn, act = g[0], int(g[1]), int(g[2]), int(g[4])
    celda(ws, fila, 1, nombre)
    celda(ws, fila, 2, soc, '#,##0')
    celda(ws, fila, 3, pn, '#,##0')
    celda(ws, fila, 4, f'=B{fila}+C{fila}', '#,##0')
    celda(ws, fila, 5, act, '#,##0')
    FILAS_SRI[nombre.split(' ')[0]] = fila
    fila += 1
ult = fila - 1
celda(ws, fila, 1, 'Total', negrita=True)
for col in range(2, 6):
    L = get_column_letter(col)
    celda(ws, fila, col, f'=SUM({L}5:{L}{ult})', '#,##0', negrita=True)
FILA_TOT_SRI = fila
fila += 2
ws.cell(row=fila, column=1, value='Empresas formales de las secciones O, P, Q y U (excluidas del universo "actividades productivas" del INEC)').font = NEGRITA
fila += 1
celda(ws, fila, 1, 'Administración pública (O), Enseñanza (P), Salud (Q) y otras (U)')
celda(ws, fila, 4, 416, '#,##0')
FILA_OPQU = fila
ws.cell(row=fila, column=5, value='O = 96, P = 119, Q = 201, U = 0 (cálculo propio sobre el catastro)').font = Font(italic=True, size=9)
fila += 2
ws.cell(row=fila, column=1, value='Empresas formales de los grupos objetivo por cantón').font = NEGRITA
fila += 1
cab(ws, fila, ['Cantón', 'Empresas formales en grupos G1 a G9'])
fila += 1
for r in leer_csv('sri_cantones.csv')[1:]:
    celda(ws, fila, 1, r[0].title())
    celda(ws, fila, 2, int(r[1]), '#,##0')
    fila += 1
fila += 1
ws.cell(row=fila, column=1, value='Altas de sociedades por año de inicio de actividades (RUC hoy activos)').font = NEGRITA
fila += 1
cab(ws, fila, ['Año', 'Sociedades activas que iniciaron ese año'])
fila += 1
for r in leer_csv('sri_altas.csv')[1:]:
    if int(r[0]) >= 2015:
        celda(ws, fila, 1, int(r[0]) if int(r[0]) < 2026 else '2026 (parcial)')
        celda(ws, fila, 2, int(r[1]), '#,##0')
        fila += 1
ws.cell(row=fila + 1, column=1, value='Nota: incluye solo RUC que siguen activos, por lo que los años antiguos están subestimados.').font = Font(italic=True, size=9)
anchos(ws, [62, 20, 26, 20, 26])

# ------------------------------------------------------------------ INEC
ws = wb.create_sheet('INEC_LosRios')
titulo(ws, 'Indicadores del INEC para Los Ríos', 'Fuentes: REEM 2025, ENEMDU anual 2025, Censo de Población 2022.')
cab(ws, 4, ['Indicador', 'Los Ríos', 'Nacional', 'Fuente'])
inec = [
    ('Empresas del universo REEM, 2025', 32993, 1204276, 'REEM 2025, cuadro 1.1.8'),
    ('Empresas con actividades productivas (ventas y empleo IESS), 2025', 1968, 86311, 'REEM 2025, cuadro 2.1.7'),
    ('Empleo registrado equivalente (IESS), 2025', 66742.5, 2921188.6, 'REEM 2025, cuadro 1.3.18'),
    ('Empleo adecuado o pleno, 2025 (%)', 32.5, 37.1, 'ENEMDU anual 2025, tabla 5 y sección 1.4.1'),
    ('Subempleo, 2025 (%)', 28.4, None, 'ENEMDU anual 2025, tabla 6'),
    ('Desempleo, 2025 (%)', 1.7, 3.6, 'ENEMDU anual 2025, tabla 3 y sección 1.1'),
    ('Población, Censo 2022', 898652, None, 'INEC, Censo 2022 (Los Ríos, cuarta provincia más poblada)'),
]
for i, (k, a, b, f) in enumerate(inec, start=5):
    celda(ws, i, 1, k)
    celda(ws, i, 2, a, '#,##0.0' if isinstance(a, float) else '#,##0')
    celda(ws, i, 3, b, '#,##0.0' if isinstance(b, float) else '#,##0')
    celda(ws, i, 4, f)
FILA_EMP_PROD = 6
FILA_POB = 11
anchos(ws, [64, 16, 16, 52])

# ------------------------------------------------------------------ Supuestos
ws = wb.create_sheet('Supuestos')
titulo(ws, 'Supuestos del modelo (celdas amarillas: editables)',
       'Sin fuente oficial. Son valores de trabajo para dimensionar; sustitúyalos con los resultados de la encuesta a empleadores.')

ws['A4'] = '1. Factor de empresas operativas'
ws['A4'].font = NEGRITA
celda(ws, 5, 1, 'Empresas INEC con ventas y empleo IESS (Los Ríos, 2025)')
celda(ws, 5, 2, f"=INEC_LosRios!B{FILA_EMP_PROD}", '#,##0')
celda(ws, 6, 1, 'Empresas formales SRI sin secciones O, P, Q y U')
celda(ws, 6, 2, f"=SRI_LosRios!D{FILA_TOT_SRI}-SRI_LosRios!D{FILA_OPQU}", '#,##0')
celda(ws, 7, 1, 'Factor operativo calculado (INEC / SRI)', negrita=True)
celda(ws, 7, 2, '=B5/B6', '0.0%')
celda(ws, 8, 1, 'Factor operativo que usa el modelo (por defecto, el calculado, redondeado)')
celda(ws, 8, 2, '=ROUND(B7,2)', '0.0%', entrada=True)
celda(ws, 9, 1, 'Factor operativo del sector público (los GAD operan)')
celda(ws, 9, 2, 1.0, '0.0%', entrada=True)
ws['C7'] = 'Aproximado: las definiciones de INEC y SRI no coinciden exactamente.'
ws['C7'].font = Font(italic=True, size=9)

ws['A11'] = '2. Adopción: % de empresas operativas que contratan al menos un perfil IoT'
ws['A11'].font = NEGRITA
cab(ws, 12, ['Grupo', 'Conservador 2030', 'Conservador 2031', 'Base 2030', 'Base 2031', 'Optimista 2030', 'Optimista 2031'])
adop = leer_csv('supuestos_adopcion.csv')[1:]
fila = 13
FILA_ADOP = {}
for r in adop:
    celda(ws, fila, 1, r[0])
    for j in range(6):
        celda(ws, fila, 2 + j, float(r[1 + j]) / 100, '0.0%', entrada=True)
    FILA_ADOP[r[0].split(' ')[0]] = fila
    fila += 1
ws.cell(row=fila, column=1, value='Los porcentajes de la hoja de adopción son supuestos del equipo, sin fuente oficial.').font = Font(italic=True, size=9)

fila += 2
ws.cell(row=fila, column=1, value='3. Profesionales IoT por empresa que adopta').font = NEGRITA
fila += 1
cab(ws, fila, ['', 'Conservador', 'Base', 'Optimista'])
fila += 1
celda(ws, fila, 1, 'Profesionales por empresa adoptante')
for j, v in enumerate([1.0, 1.5, 2.0]):
    celda(ws, fila, 2 + j, v, '0.0', entrada=True)
FILA_PROF = fila
fila += 1
celda(ws, fila, 1, 'Profesionales IoT que ya trabajan en las empresas en 2026')
celda(ws, fila, 2, 0, '#,##0', entrada=True)
FILA_EXIST = fila

fila += 2
ws.cell(row=fila, column=1, value='4. Oferta de graduados TIC en Los Ríos').font = NEGRITA
fila += 1
celda(ws, fila, 1, 'Graduados TIC en el país, 2021 (SENESCYT, vía MINTEL)')
celda(ws, fila, 2, 7476, '#,##0')
FILA_NAC = fila
fila += 1
celda(ws, fila, 1, 'Graduados TIC en el país, 2020 (SENESCYT, vía MINTEL)')
celda(ws, fila, 2, 5079, '#,##0')
fila += 1
celda(ws, fila, 1, 'Población de Los Ríos sobre la nacional (proxy máximo)')
celda(ws, fila, 2, '=INEC_LosRios!B%d/16938986' % FILA_POB, '0.0%')
ws.cell(row=fila, column=3, value='Población nacional del Censo 2022 tomada como 16.938.986; verificar').font = Font(italic=True, size=9)
FILA_SHARE = fila
fila += 1
cab(ws, fila, ['', 'Oferta baja', 'Oferta base', 'Oferta alta'])
fila += 1
celda(ws, fila, 1, 'Graduados TIC por año que salen de IES de Los Ríos')
celda(ws, fila, 2, 69, '#,##0', entrada=True)
celda(ws, fila, 3, f'=ROUND(B{FILA_NAC}*0.025,0)', '#,##0', entrada=True)
celda(ws, fila, 4, f'=ROUND(B{FILA_NAC}*B{FILA_SHARE},0)', '#,##0', entrada=True)
FILA_GRAD = fila
ws.cell(row=fila + 1, column=1, value='Baja: máximo identificado en las carreras TIC de la UTEQ en la cohorte 2018-2019 (45 a 69 graduados). Base: 2,5 % del total nacional. Alta: participación de la población de Los Ríos (5,3 %).').font = Font(italic=True, size=9)
fila += 2
celda(ws, fila, 1, '% de esos graduados con competencias útiles para IoT')
celda(ws, fila, 2, 0.02, '0.0%', entrada=True)
celda(ws, fila, 3, 0.05, '0.0%', entrada=True)
celda(ws, fila, 4, 0.10, '0.0%', entrada=True)
FILA_PIOT = fila

fila += 2
ws.cell(row=fila, column=1, value='5. La carrera propuesta').font = NEGRITA
fila += 1
celda(ws, fila, 1, 'Estudiantes por cohorte')
celda(ws, fila, 2, 40, '#,##0', entrada=True)
FILA_COH = fila
fila += 1
celda(ws, fila, 1, 'Tasa de titulación (UTEQ, cohorte 2018-2019)')
celda(ws, fila, 2, 0.6096, '0.0%', entrada=True)
FILA_TIT = fila
fila += 1
celda(ws, fila, 1, 'Año del primer ingreso')
celda(ws, fila, 2, 2027, '0', entrada=True)
FILA_ING = fila
fila += 1
celda(ws, fila, 1, 'Duración de la carrera en años')
celda(ws, fila, 2, 4, '0', entrada=True)
FILA_DUR = fila
anchos(ws, [78, 20, 20, 20, 20, 20, 20])

# ------------------------------------------------------------------ Demanda
ws = wb.create_sheet('Demanda')
titulo(ws, 'Demanda estimada de profesionales IoT en Los Ríos',
       'Puestos = empresas formales × factor operativo × % que adopta × profesionales por empresa. Es un stock acumulado al año indicado.')
cab(ws, 4, ['Grupo', 'Empresas formales (SRI)', 'Factor operativo', 'Empresas operativas',
            'Adoptan Cons. 2030', 'Adoptan Cons. 2031', 'Adoptan Base 2030', 'Adoptan Base 2031', 'Adoptan Opt. 2030', 'Adoptan Opt. 2031'])
fila = 5
claves = list(FILA_ADOP.keys())
ini = fila
for k in claves:
    fs = FILAS_SRI[k]
    fa = FILA_ADOP[k]
    celda(ws, fila, 1, f'=SRI_LosRios!A{fs}')
    celda(ws, fila, 2, f'=SRI_LosRios!D{fs}', '#,##0')
    celda(ws, fila, 3, '=Supuestos!$B$9' if k == 'G8' else '=Supuestos!$B$8', '0%')
    celda(ws, fila, 4, f'=B{fila}*C{fila}', '#,##0.0')
    for j in range(6):
        col_sup = get_column_letter(2 + j)
        celda(ws, fila, 5 + j, f'=D{fila}*Supuestos!{col_sup}{fa}', '#,##0.0')
    fila += 1
fin = fila - 1
celda(ws, fila, 1, 'Total empresas que adoptan', negrita=True)
for col in [2, 4, 5, 6, 7, 8, 9, 10]:
    L = get_column_letter(col)
    celda(ws, fila, col, f'=SUM({L}{ini}:{L}{fin})', '#,##0.0', negrita=True)
FILA_TOT_ADOP = fila
fila += 2
ws.cell(row=fila, column=1, value='Puestos de perfil IoT (stock acumulado)').font = NEGRITA
fila += 1
cab(ws, fila, ['Escenario', 'A 4 años (2030)', 'A 5 años (2031)', 'Promedio anual nuevo 2027-2031'])
fila += 1
FILA_DEM = {}
for nombre, c4, c5, colprof in [('Conservador', 'E', 'F', 'B'), ('Base', 'G', 'H', 'C'), ('Optimista', 'I', 'J', 'D')]:
    celda(ws, fila, 1, nombre, negrita=True)
    celda(ws, fila, 2, f'={c4}{FILA_TOT_ADOP}*Supuestos!{colprof}{FILA_PROF}-Supuestos!$B${FILA_EXIST}', '#,##0')
    celda(ws, fila, 3, f'={c5}{FILA_TOT_ADOP}*Supuestos!{colprof}{FILA_PROF}-Supuestos!$B${FILA_EXIST}', '#,##0')
    celda(ws, fila, 4, f'=C{fila}/5', '#,##0.0')
    FILA_DEM[nombre] = fila
    fila += 1
ws.cell(row=fila + 1, column=1, value='Por grupo, escenario base 2031 (puestos)').font = NEGRITA
fila += 2
cab(ws, fila, ['Grupo', 'Puestos base 2031', '% del total'])
fila += 1
ini2 = fila
for i, k in enumerate(claves):
    celda(ws, fila, 1, f'=A{ini + i}')
    celda(ws, fila, 2, f'=H{ini + i}*Supuestos!$C${FILA_PROF}', '#,##0.0')
    celda(ws, fila, 3, f'=B{fila}/SUM($B${ini2}:$B${ini2 + len(claves) - 1})', '0.0%')
    fila += 1
anchos(ws, [56, 18, 14, 18, 18, 18, 18, 18, 18, 18])

# ------------------------------------------------------------------ Oferta
ws = wb.create_sheet('Oferta')
titulo(ws, 'Oferta de profesionales con perfil IoT en Los Ríos',
       'Sin la carrera: graduados TIC de IES de Los Ríos con competencias IoT. Con la carrera: se suma su promoción desde el primer año de graduación.')
cab(ws, 4, ['Concepto', 'Oferta baja', 'Oferta base', 'Oferta alta'])
celda(ws, 5, 1, 'Graduados TIC por año en IES de Los Ríos')
celda(ws, 6, 1, '% con competencias útiles para IoT')
celda(ws, 7, 1, 'Graduados con perfil IoT por año (sin la carrera)', negrita=True)
for j, L in enumerate(['B', 'C', 'D']):
    celda(ws, 5, 2 + j, f'=Supuestos!{L}{FILA_GRAD}', '#,##0')
    celda(ws, 6, 2 + j, f'=Supuestos!{L}{FILA_PIOT}', '0.0%')
    celda(ws, 7, 2 + j, f'={L}5*{L}6', '#,##0.0', negrita=True)
ws['A9'] = 'Acumulado de graduados con perfil IoT sin la carrera'
ws['A9'].font = NEGRITA
cab(ws, 10, ['Período', 'Oferta baja', 'Oferta base', 'Oferta alta'])
celda(ws, 11, 1, '2027-2030 (4 años)')
celda(ws, 12, 1, '2027-2031 (5 años)')
for j, L in enumerate(['B', 'C', 'D']):
    celda(ws, 11, 2 + j, f'={L}7*4', '#,##0.0')
    celda(ws, 12, 2 + j, f'={L}7*5', '#,##0.0')
ws['A14'] = 'Graduados de la carrera propuesta'
ws['A14'].font = NEGRITA
cab(ws, 15, ['Concepto', 'Valor'])
celda(ws, 16, 1, 'Primer año de graduación')
celda(ws, 16, 2, f'=Supuestos!B{FILA_ING}+Supuestos!B{FILA_DUR}', '0')
celda(ws, 17, 1, 'Graduados por promoción')
celda(ws, 17, 2, f'=Supuestos!B{FILA_COH}*Supuestos!B{FILA_TIT}', '#,##0.0')
celda(ws, 18, 1, 'Graduados acumulados hasta 2030')
celda(ws, 18, 2, '=MAX(0,2030-B16+1)*B17', '#,##0.0')
celda(ws, 19, 1, 'Graduados acumulados hasta 2031')
celda(ws, 19, 2, '=MAX(0,2031-B16+1)*B17', '#,##0.0')
celda(ws, 20, 1, 'Graduados por año desde el segundo año de graduación')
celda(ws, 20, 2, '=B17', '#,##0.0')
anchos(ws, [60, 18, 18, 18])

# ------------------------------------------------------------------ Brecha
ws = wb.create_sheet('Brecha')
titulo(ws, 'Brecha entre demanda y oferta de perfiles IoT en Los Ríos',
       'Brecha = puestos demandados − graduados con perfil IoT disponibles. Positiva: faltan profesionales.')
d = 'Demanda'
FILAS_B = {}
r = 4
for horizonte, colD, fila_of, col_car in [('A 4 años (2030)', 'B', 11, 'B18'), ('A 5 años (2031)', 'C', 12, 'B19')]:
    ws.cell(row=r, column=1, value=f'Horizonte {horizonte}').font = Font(bold=True, size=12, color=AZUL)
    r += 1
    ws.cell(row=r, column=1, value='Sin la carrera').font = NEGRITA
    r += 1
    cab(ws, r, ['Demanda \\ Oferta', 'Oferta baja', 'Oferta base', 'Oferta alta'])
    r += 1
    for esc in ['Conservador', 'Base', 'Optimista']:
        celda(ws, r, 1, f'Demanda {esc.lower()}', negrita=True)
        for j, L in enumerate(['B', 'C', 'D']):
            celda(ws, r, 2 + j, f'={d}!{colD}{FILA_DEM[esc]}-Oferta!{L}{fila_of}', '#,##0;[Red]-#,##0')
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='Con la carrera').font = NEGRITA
    r += 1
    cab(ws, r, ['Demanda \\ Oferta', 'Oferta baja', 'Oferta base', 'Oferta alta'])
    r += 1
    for esc in ['Conservador', 'Base', 'Optimista']:
        celda(ws, r, 1, f'Demanda {esc.lower()}', negrita=True)
        for j, L in enumerate(['B', 'C', 'D']):
            celda(ws, r, 2 + j, f'={d}!{colD}{FILA_DEM[esc]}-Oferta!{L}{fila_of}-Oferta!{col_car}', '#,##0;[Red]-#,##0')
        FILAS_B[(horizonte, esc)] = r
        r += 1
    r += 2

ws.cell(row=r, column=1, value='Flujo anual en régimen (desde 2032, con la carrera funcionando)').font = Font(bold=True, size=12, color=AZUL)
r += 1
cab(ws, r, ['Escenario', 'Puestos nuevos por año', 'Graduados IoT por año sin la carrera (oferta base)', 'Graduados por año de la carrera', 'Cobertura sin la carrera', 'Cobertura con la carrera'])
r += 1
ini_f = r
for esc in ['Conservador', 'Base', 'Optimista']:
    celda(ws, r, 1, esc, negrita=True)
    celda(ws, r, 2, f'={d}!D{FILA_DEM[esc]}', '#,##0.0')
    celda(ws, r, 3, '=Oferta!C7', '#,##0.0')
    celda(ws, r, 4, '=Oferta!B20', '#,##0.0')
    celda(ws, r, 5, f'=C{r}/B{r}', '0%')
    celda(ws, r, 6, f'=(C{r}+D{r})/B{r}', '0%')
    r += 1
ws.cell(row=r, column=1, value='Cobertura mayor a 100 %: la oferta supera los puestos nuevos, lo que sería sobreoferta.').font = Font(italic=True, size=9)
fin_f = r - 1
ch = BarChart()
ch.type = 'col'
ch.title = 'Puestos nuevos por año y graduados IoT por año'
ch.y_axis.title = 'Personas por año'
data = Reference(ws, min_col=2, max_col=4, min_row=ini_f - 1, max_row=fin_f)
cats = Reference(ws, min_col=1, min_row=ini_f, max_row=fin_f)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.height, ch.width = 8, 18
ws.add_chart(ch, f'A{r + 3}')
anchos(ws, [34, 22, 30, 24, 22, 22])

# ------------------------------------------------------------------ Sensibilidad
ws = wb.create_sheet('Sensibilidad')
titulo(ws, 'Sensibilidad de la demanda a 5 años (2031)',
       'Puestos IoT si la adopción base se multiplica por un factor y cambia el número de profesionales por empresa adoptante.')
celda(ws, 4, 1, 'Empresas que adoptan en 2031, escenario base')
celda(ws, 4, 2, f'=Demanda!H{FILA_TOT_ADOP}', '#,##0.0')
cab(ws, 6, ['Multiplicador de adopción \\ profesionales por empresa', 1.0, 1.5, 2.0, 3.0])
for j, v in enumerate([1.0, 1.5, 2.0, 3.0]):
    ws.cell(row=6, column=2 + j).number_format = '0.0'
for i, m in enumerate([0.25, 0.5, 1.0, 1.5, 2.0]):
    rr = 7 + i
    celda(ws, rr, 1, m, '0.00"x"', entrada=True)
    for j in range(4):
        L = get_column_letter(2 + j)
        celda(ws, rr, 2 + j, f'=$B$4*$A{rr}*{L}$6', '#,##0')
ws['A13'] = 'Lectura: la fila 1,00x y la columna 1,5 reproducen el escenario base. Compare con los graduados por año de la hoja Brecha.'
ws['A13'].font = Font(italic=True, size=9)
anchos(ws, [56, 14, 14, 14, 14])

wb.save(SALIDA)
print('OK', SALIDA)
